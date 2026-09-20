#!/usr/bin/env python3
"""Create an ITSME-style Excel summary from ITSx classification CSVs.

The input directory should contain one CSV per sample, produced by the ITSx
classification workflow. Each CSV must contain ``consensus_taxonomy`` and the
five ITSx coordinate columns.

Examples
--------
With expected taxonomy and rank-match columns::

    python summary.itsx.py itsx_summaries \
        Arctic_taxonomy_expectations_updated.csv \
        -o ITSx_master_taxonomy_summary.xlsx

Without an expectation table::

    python summary.itsx.py itsx_summaries \
        -o ITSx_master_taxonomy_summary.xlsx

Dependency
----------
    python -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Install it with: python -m pip install openpyxl"
    ) from exc


SEPARATOR = " || "
RANKS = [
    "domain",
    "kingdom",
    "phylum",
    "class",
    "order",
    "family",
    "genus",
    "species",
]
REGION_COLUMNS = [
    ("SSU_coordinates", "18S"),
    ("ITS1_coordinates", "ITS1"),
    ("5.8S_coordinates", "5.8S"),
    ("ITS2_coordinates", "ITS2"),
    ("LSU_coordinates", "28S"),
]
FULL_LOCUS_COLUMNS = [column for column, _ in REGION_COLUMNS]
MISSING_VALUES = {
    "",
    "na",
    "n/a",
    "nan",
    "none",
    "not found",
    "unknown",
}
PARTIAL_COORDINATES = {"no start", "no end"}
COORDINATE_RE = re.compile(r"^\s*\d+\s*-\s*\d+\s*$")
METAZOAN_PHYLA = {
    "annelida",
    "arthropoda",
    "brachiopoda",
    "bryozoa",
    "chaetognatha",
    "chordata",
    "cnidaria",
    "ctenophora",
    "echinodermata",
    "gastrotricha",
    "hemichordata",
    "mollusca",
    "nematoda",
    "nemertea",
    "onychophora",
    "platyhelminthes",
    "porifera",
    "rotifera",
    "tardigrada",
}
RANK_COLUMN_ALIASES = {
    "domain": ["Domain", "domain", "Superkingdom", "superkingdom"],
    "kingdom": ["Kingdom", "kingdom"],
    "phylum": ["Phylum", "phylum"],
    "class": ["Class", "class"],
    "order": ["Order", "order"],
    "family": ["Family", "family"],
    "genus": ["Genus", "genus"],
    "species": ["Species/Taxa", "Species", "species", "Taxon", "taxon"],
}
EXPECTED_TAXONOMY_COLUMNS = [
    "Taxa_String",
    "taxa_string",
    "Expected taxonomy",
    "expected_taxonomy",
]
SAMPLE_COLUMNS = ["sample", "Sample", "sample_id", "Sample ID"]


def clean(value: object) -> str:
    """Return a stripped string, converting None to an empty string."""
    return "" if value is None else str(value).strip()


def normalized(value: object) -> str:
    """Normalize taxonomy labels for case-insensitive exact matching."""
    return " ".join(clean(value).lower().split())


def natural_key(value: str) -> list[object]:
    """Sort identifiers naturally, for example B2 before B10."""
    return [
        int(piece) if piece.isdigit() else piece.lower()
        for piece in re.split(r"(\d+)", value)
    ]


def first_present(row: dict[str, str], names: list[str]) -> str:
    for name in names:
        value = clean(row.get(name))
        if value:
            return value
    return ""


def taxonomy_tokens(value: object) -> list[str]:
    return [normalized(token) for token in clean(value).split(";") if clean(token)]


def read_expected_taxonomies(path: Path | None) -> dict[str, list[dict[str, object]]]:
    """Read one or more expected lineages for each sample.

    The table may provide ``Taxa_String`` (normally phylum through species),
    rank-specific columns, or both. Animal expectations that begin at phylum
    are assigned Eukaryota and Metazoa for the two omitted upper ranks.
    """
    if path is None:
        return {}
    if not path.is_file():
        raise FileNotFoundError(f"Expected-taxonomy CSV not found: {path}")

    result: defaultdict[str, list[dict[str, object]]] = defaultdict(list)
    seen: defaultdict[str, set[str]] = defaultdict(set)

    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        sample_column = next((name for name in SAMPLE_COLUMNS if name in fields), None)
        taxonomy_column = next(
            (name for name in EXPECTED_TAXONOMY_COLUMNS if name in fields), None
        )
        rank_columns_present = any(
            alias in fields for aliases in RANK_COLUMN_ALIASES.values() for alias in aliases
        )
        if sample_column is None:
            raise ValueError(
                f"{path} needs a sample column (accepted names: "
                f"{', '.join(SAMPLE_COLUMNS)})"
            )
        if taxonomy_column is None and not rank_columns_present:
            raise ValueError(
                f"{path} needs Taxa_String/expected_taxonomy or rank columns"
            )

        for row in reader:
            sample = clean(row.get(sample_column))
            if not sample:
                continue

            ranks: dict[str, str] = {}
            for rank in RANKS:
                ranks[rank] = first_present(row, RANK_COLUMN_ALIASES[rank])

            taxonomy = clean(row.get(taxonomy_column)) if taxonomy_column else ""
            tokens = [clean(token) for token in taxonomy.split(";") if clean(token)]
            if tokens:
                string_ranks = RANKS if len(tokens) >= len(RANKS) else RANKS[2:]
                for rank, token in zip(string_ranks, tokens):
                    if not ranks[rank]:
                        ranks[rank] = token

            if ranks["phylum"] and not ranks["domain"]:
                ranks["domain"] = "Eukaryota"
            if normalized(ranks["phylum"]) in METAZOAN_PHYLA and not ranks["kingdom"]:
                ranks["kingdom"] = "Metazoa"

            if not taxonomy:
                taxonomy = ";".join(ranks[rank] for rank in RANKS[2:] if ranks[rank])
            if not taxonomy:
                continue

            lineage_key = normalized(taxonomy)
            if lineage_key in seen[sample]:
                continue
            seen[sample].add(lineage_key)
            result[sample].append({"taxonomy": taxonomy, "ranks": ranks})

    return dict(result)


def sample_from_filename(path: Path, expected_samples: set[str]) -> str:
    """Derive a sample name from a per-sample ITSx summary filename."""
    sample = path.stem.strip()

    for prefix in ("itsme_", "itsx_"):
        if sample.lower().startswith(prefix):
            sample = sample[len(prefix) :]
            break

    suffixes = (
        ".master_summary",
        "_master_summary",
        "_rrna_analysis_sensitive",
        "_rrna_analysis",
        "_spades",
        "-spades",
        "_itsx",
    )
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if sample.lower().endswith(suffix.lower()):
                sample = sample[: -len(suffix)]
                changed = True
                break

    # Finder/download duplicates such as SAMPLE-2 are repaired only if the
    # unsuffixed identifier is present in the expectation table.
    base = re.sub(r"-\d+$", "", sample)
    if base != sample and base in expected_samples:
        sample = base

    return sample


def coordinate_state(value: object) -> str:
    """Return complete, partial, or absent for an ITSx coordinate value."""
    text = clean(value).lower()
    if COORDINATE_RE.fullmatch(text):
        return "complete"
    if text in PARTIAL_COORDINATES:
        return "partial"
    return "absent"


def its_type(row: dict[str, str]) -> str:
    """Call Full only for a complete 18S-ITS1-5.8S-ITS2-28S layout."""
    if all(
        coordinate_state(row.get(column)) == "complete"
        for column in FULL_LOCUS_COLUMNS
    ):
        return "Full"
    return "Partial"


def recovered_regions(row: dict[str, str]) -> str:
    """List detected rRNA/ITS components in genomic order."""
    parts: list[str] = []
    for column, label in REGION_COLUMNS:
        state = coordinate_state(row.get(column))
        if state == "complete":
            parts.append(label)
        elif state == "partial":
            parts.append(f"{label} (partial)")

    if not parts:
        locus_type = clean(row.get("locus_type")).upper()
        if "18S" in locus_type or "SSU" in locus_type:
            parts.append("18S")
        if "ITS1" in locus_type:
            parts.append("ITS1")
        if "5.8S" in locus_type or "5_8S" in locus_type:
            parts.append("5.8S")
        if "ITS2" in locus_type:
            parts.append("ITS2")
        if "28S" in locus_type or "LSU" in locus_type:
            parts.append("28S")

    return ", ".join(parts) if parts else "Unspecified"


def rank_match_scores(
    predicted_taxonomy: str,
    expected_records: list[dict[str, object]],
) -> dict[str, int | None]:
    """Score every rank independently against any expected lineage."""
    if not expected_records:
        return {rank: None for rank in RANKS}

    predicted_tokens = taxonomy_tokens(predicted_taxonomy)
    predicted = dict(zip(RANKS, predicted_tokens))
    scores: dict[str, int | None] = {}

    for rank in RANKS:
        expected_values = {
            normalized(record["ranks"].get(rank))
            for record in expected_records
            if clean(record["ranks"].get(rank))
        }
        scores[rank] = int(
            bool(predicted.get(rank)) and predicted[rank] in expected_values
        )

    return scores


def read_itsx_summaries(
    folder: Path,
    expected_samples: set[str],
    include_unresolved: bool,
) -> tuple[list[dict[str, str]], int]:
    """Read classified ITSx CSVs and collapse duplicate summary calls."""
    if not folder.is_dir():
        raise NotADirectoryError(f"Not a directory: {folder}")

    paths = sorted(folder.glob("*.csv"), key=lambda item: natural_key(item.name))
    if not paths:
        raise FileNotFoundError(f"No CSV files found in {folder}")

    calls: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()
    skipped_unresolved = 0

    for path in paths:
        sample = sample_from_filename(path, expected_samples)
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            fields = set(reader.fieldnames or [])
            required = {"consensus_taxonomy", *[column for column, _ in REGION_COLUMNS]}
            missing = sorted(required - fields)
            if missing:
                raise ValueError(
                    f"{path} is not an ITSx classification summary; missing: "
                    f"{', '.join(missing)}"
                )

            for row in reader:
                taxonomy = clean(row.get("consensus_taxonomy"))
                unresolved = normalized(taxonomy) in MISSING_VALUES
                if unresolved and not include_unresolved:
                    skipped_unresolved += 1
                    continue
                if unresolved:
                    taxonomy = "Unresolved"

                call_type = its_type(row)
                regions = recovered_regions(row)
                key = (sample, call_type, normalized(taxonomy), regions)
                if key in seen:
                    continue
                seen.add(key)
                calls.append(
                    {
                        "sample": sample,
                        "type": call_type,
                        "taxonomy": taxonomy,
                        "regions": regions,
                    }
                )

    return calls, skipped_unresolved


def create_workbook(
    calls: list[dict[str, str]],
    expected: dict[str, list[dict[str, object]]],
    output_path: Path,
) -> tuple[int, int, int]:
    """Write the formatted ITSME-style workbook."""
    calls.sort(
        key=lambda call: (
            natural_key(call["sample"]),
            0 if call["type"] == "Full" else 1,
            call["taxonomy"].lower(),
            call["regions"],
        )
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ITS calls"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "1F4E78"

    samples = sorted({call["sample"] for call in calls}, key=natural_key)
    full_count = sum(call["type"] == "Full" for call in calls)
    partial_count = sum(call["type"] == "Partial" for call in calls)

    worksheet["A1"] = "ITSx ITS taxonomy calls"
    worksheet["A1"].font = Font(name="Arial", size=14, bold=True, color="1F2937")
    worksheet["A2"] = (
        f"{len(samples)} samples; {full_count} full calls; "
        f"{partial_count} partial calls"
    )
    worksheet["A2"].font = Font(name="Arial", size=10, italic=True, color="5B6573")

    if expected:
        note = (
            "Rank scoring: 1 = exact match to at least one expected taxon; "
            "0 = mismatch, unresolved rank, or missing expected rank."
        )
    else:
        note = "No expected-taxonomy table supplied; expected taxonomy and rank matches are blank."
    worksheet["A3"] = note
    worksheet["A3"].font = Font(name="Arial", size=10, italic=True, color="5B6573")

    header_row = 4
    headers = [
        "Sample",
        "ITS type",
        "ITS taxonomy",
        "Recovered regions",
        "Expected taxonomy",
        "Domain match",
        "Kingdom match",
        "Phylum match",
        "Class match",
        "Order match",
        "Family match",
        "Genus match",
        "Species match",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column, header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    light_blue = PatternFill("solid", fgColor="F3F7FB")
    white = PatternFill("solid", fgColor="FFFFFF")
    full_fill = PatternFill("solid", fgColor="DDF4E4")
    partial_fill = PatternFill("solid", fgColor="FFF0CC")
    top_side = Side(style="medium", color="6B7A90")
    thin_side = Side(style="hair", color="D8E0E8")

    current_sample: str | None = None
    sample_index = -1
    excel_row = header_row + 1

    for call in calls:
        sample = call["sample"]
        expected_records = expected.get(sample, [])
        expected_text = SEPARATOR.join(
            clean(record["taxonomy"]) for record in expected_records
        )
        scores = rank_match_scores(call["taxonomy"], expected_records)

        is_new_sample = sample != current_sample
        if is_new_sample:
            current_sample = sample
            sample_index += 1

        values: list[object] = [
            sample,
            call["type"],
            call["taxonomy"],
            call["regions"],
            expected_text,
            *[scores[rank] for rank in RANKS],
        ]

        block_fill = light_blue if sample_index % 2 == 0 else white
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(excel_row, column, value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.fill = block_fill
            cell.alignment = Alignment(
                vertical="center",
                horizontal=(
                    "center"
                    if column in (2, 4, 6, 7, 8, 9, 10, 11, 12, 13)
                    else "left"
                ),
                wrap_text=column in (3, 4, 5),
            )
            cell.border = Border(top=top_side if is_new_sample else thin_side)

        worksheet.cell(excel_row, 1).font = Font(
            name="Arial", size=10, bold=True, color="243B53"
        )
        type_cell = worksheet.cell(excel_row, 2)
        if call["type"] == "Full":
            type_cell.fill = full_fill
            type_cell.font = Font(
                name="Arial", size=10, bold=True, color="176A3A"
            )
        else:
            type_cell.fill = partial_fill
            type_cell.font = Font(
                name="Arial", size=10, bold=True, color="8A5A00"
            )

        excel_row += 1

    last_row = max(header_row, excel_row - 1)
    worksheet.auto_filter.ref = f"A{header_row}:M{last_row}"
    worksheet.freeze_panes = "B5"
    worksheet.row_dimensions[header_row].height = 26

    widths = {
        "A": 25,
        "B": 11,
        "C": 72,
        "D": 28,
        "E": 66,
        "F": 13,
        "G": 13,
        "H": 12,
        "I": 12,
        "J": 12,
        "K": 13,
        "L": 12,
        "M": 13,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    for row_number in range(header_row + 1, last_row + 1):
        worksheet.row_dimensions[row_number].height = 30

    worksheet.print_title_rows = f"1:{header_row}"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(samples), full_count, partial_count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create a block-formatted, long-format Excel summary from a folder "
            "of ITSx classification CSVs."
        )
    )
    parser.add_argument(
        "itsx_summaries_dir",
        type=Path,
        help="Directory containing per-sample ITSx classification CSVs",
    )
    parser.add_argument(
        "expected_csv",
        nargs="?",
        type=Path,
        help=(
            "Optional expected-taxonomy CSV; needed to populate Expected "
            "taxonomy and rank-match columns"
        ),
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("ITSx_master_taxonomy_summary.xlsx"),
        help="Output Excel workbook (default: ITSx_master_taxonomy_summary.xlsx)",
    )
    parser.add_argument(
        "--include-unresolved",
        action="store_true",
        help="Include rows whose consensus_taxonomy is NA as 'Unresolved'",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.output.suffix.lower() != ".xlsx":
        raise SystemExit("The output filename must end in .xlsx")

    try:
        expected = read_expected_taxonomies(args.expected_csv)
        calls, skipped_unresolved = read_itsx_summaries(
            args.itsx_summaries_dir,
            set(expected),
            args.include_unresolved,
        )
        sample_count, full_count, partial_count = create_workbook(
            calls, expected, args.output
        )
    except (FileNotFoundError, NotADirectoryError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    message = (
        f"Wrote {len(calls)} ITS calls from {sample_count} samples "
        f"({full_count} full, {partial_count} partial) to {args.output}"
    )
    if skipped_unresolved and not args.include_unresolved:
        message += f"; skipped {skipped_unresolved} unresolved rows"
    print(message)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
